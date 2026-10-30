import sqlite3
from collections import Counter

from openpyxl import Workbook

from . import storage
from .domain import CANDIDATE, DEAD_SITE, GUESSED_SITE, NO_SITE, ONLY_SOCIAL

HEAD = ["Nome", "Categoria", "Indirizzo", "Telefono", "Zona", "Esito", "Dettaglio", "Mappa"]
SHEETS = {
    "Senza sito (confermato)": (NO_SITE,),
    "Solo social o sito morto": (ONLY_SOCIAL, DEAD_SITE),
    "Da verificare": (CANDIDATE, GUESSED_SITE),
}


def _line(r: sqlite3.Row) -> list[str]:
    maps = f"https://www.google.com/maps/search/?api=1&query={r['lat']},{r['lon']}"
    return [r["name"], r["category"], r["address"], r["phone"], r["area"], r["verdict"] or "",
            r["detail"] or "", maps]


def write(con: sqlite3.Connection, path: str) -> Counter:
    rows = list(storage.all_rows(con))
    wb = Workbook()
    wb.remove(wb.active)
    for title, verdicts in SHEETS.items():
        ws = wb.create_sheet(title[:31])
        ws.append(HEAD)
        for r in rows:
            if r["verdict"] in verdicts:
                ws.append(_line(r))
        ws.freeze_panes = "A2"
    counts = Counter(r["verdict"] or "non classificato" for r in rows)
    st = wb.create_sheet("Statistiche")
    st.append(["Esito", "N"])
    for k, v in counts.most_common():
        st.append([k, v])
    st.append(["Totale", len(rows)])
    wb.save(path)
    return counts
